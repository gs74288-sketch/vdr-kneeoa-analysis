# -*- coding: utf-8 -*-
"""Internal consistency audit of the current manuscript (v11)."""
import docx, re, sys

P = r"C:\Users\liuha\Desktop\孟德尔\revision_20261002\outputs\VDR_OA_manuscript_v11_cis2supp.docx"
d = docx.Document(P)
paras = [p.text for p in d.paragraphs]
full = "\n".join(paras)

def where(s):
    return [i for i, t in enumerate(paras) if s in t]

checks = {
    "discovery MR OR 0.928":            "0.928",
    "discovery MR CI 0.869-0.992":      "0.869-0.992",
    "discovery MR P 0.027":             "P = 0.027",
    "FDR 0.438":                        "0.438",
    "10 harmonized instruments":        "10 harmonized instruments",
    "single-cell OR 1.89":              "1.89",
    "single-cell P 0.010":              "P = 0.010",
    "coloc PP.H3 0.971":                "0.971",
    "coloc PP.H4 1.9 x 10^-4":          "1.9 x 10^-4",
    "Steiger P 2.7e-145":               "2.7 x 10^-145",
    "FinnGen IVW 0.965":                "0.965",
    "FinnGen surgery 0.933":            "0.933",
    "bulk VDR P 8.5e-3":                "8.5 x 10^-3",
    "bulk NLRP3 r 0.68":                "r = 0.68",
    "bulk CYP24A1 r 0.50":              "r = 0.50",
    "GSE152805 RegC OR 1.50":           "1.50",
    "rs10783222 exclusion mentioned":   "rs10783222",
}
print("=== presence check ===")
for k, v in checks.items():
    w = where(v)
    print(f"{'OK ' if w else 'MISS'} {k:34s} occurrences={len(w)} paras={w[:6]}")

print("\n=== tables ===")
for i, t in enumerate(d.tables):
    first = t.rows[0].cells[0].text.strip()
    print(f"T{i+1}: {len(t.rows)-1} data rows x {len(t.columns)} cols | first cell: {first}")

print("\n=== instrument mentions ===")
for i, t in enumerate(paras):
    if re.search(r'\b(10|eleven|11)\b.{0,25}instrument', t, re.I):
        print(f"[{i}] {t[:200]}")

print("\n=== 'partial' / 'confirmed' wording ===")
for i, t in enumerate(paras):
    if re.search(r'partial support|not confirm|dataset-specific|not reproduced|confirmed enrichment', t, re.I):
        print(f"[{i}] {t[:170]}")
